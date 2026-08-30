# -*- coding: utf-8 -*-
"""Generate Bao_cao_tuan_9.docx and Bao_cao_tuan_10.docx (final report),
matching the established style used for tuần 4-8 (see CLAUDE.md +
scripts_tmp/gen_weekly_reports.py)."""
import os
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OLIVE = RGBColor(0x6B, 0x91, 0x1C)
BLUE_HEADER = "3B5BDB"
GRAY_CAPTION = RGBColor(0x55, 0x55, 0x55)
GRAY_FILL = "F2F2F2"


def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)


def new_doc():
    d = docx.Document()
    section = d.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2)

    normal = d.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(13)
    rpr = normal.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    h1 = d.styles['Heading 1']
    h1.font.name = 'Trebuchet MS'
    h1.font.size = Pt(20)
    h1.font.bold = True
    h1.font.color.rgb = OLIVE

    h2 = d.styles['Heading 2']
    h2.font.name = 'Trebuchet MS'
    h2.font.size = Pt(16)
    h2.font.bold = True
    h2.font.color.rgb = OLIVE

    return d


def add_cover(d, week_no):
    def p_center(text, bold=False, italic=False, size=13):
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
        r.font.size = Pt(size)
        return p

    p_center("BỘ GIÁO DỤC VÀ ĐÀO TẠO", bold=True, size=13)
    p_center("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ TP. HCM", bold=True, size=13)
    d.add_paragraph()
    p_center(f"BÁO CÁO TUẦN {week_no}", bold=True, size=20)
    p_center("ĐỒ ÁN CƠ SỞ NGÀNH CÔNG NGHỆ THÔNG TIN", bold=True, size=14)
    p_center("Đề tài: ỨNG DỤNG DI ĐỘNG HỌC TIẾNG ANH DỰA TRÊN AI", bold=True, size=13)
    p_center("Giảng viên hướng dẫn: Đặng Thị Thạch Thảo")
    p_center("Sinh viên thực hiện: Nguyễn Thị Ngọc Anh — MSSV: 2410060302")
    p_center("Trần Trí Nam — MSSV: 2410060292")
    p_center("Lớp: 24TXTHG1")
    p_center("TP. Hồ Chí Minh, 2026", italic=True)
    d.add_paragraph()


def add_heading1(d, text):
    d.add_heading(text, level=1)


def add_heading2(d, text):
    d.add_heading(text, level=2)


def add_para(d, text):
    d.add_paragraph(text)


def add_bullets(d, items):
    for it in items:
        d.add_paragraph(it, style='List Bullet')


def add_real_image(d, path, caption, height_cm=11):
    """Portrait phone screenshot — scale by height."""
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, height=Cm(height_cm))
    cap = d.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = GRAY_CAPTION
    d.add_paragraph()


def add_real_image_wide(d, path, caption, width_cm=14):
    """Landscape web screenshot (16:9) — scale by width so it stays within
    the page's content width (~16.5 cm)."""
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Cm(width_cm))
    cap = d.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = GRAY_CAPTION
    d.add_paragraph()


def add_table(d, headers, rows):
    table = d.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(hdr_cells[i], BLUE_HEADER)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(10)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            p = cells[i].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(10)
    d.add_paragraph()


ASSETS7 = r"C:\Users\Nam\Desktop\english_ai_app\screenshots_tuan7"
ASSETS9 = r"C:\Users\Nam\Desktop\english_ai_app\screenshots_tuan9"

# ---------------------------------------------------------------------------
# TUẦN 9
# ---------------------------------------------------------------------------
d9 = new_doc()
add_cover(d9, 9)

add_heading1(d9, "1. MỤC TIÊU TUẦN 9")
add_para(d9, "Báo cáo tuần 8 đã ghi nhận chức năng \"Quên mật khẩu\" là hạng mục nâng cao chưa nằm "
             "trong phạm vi ưu tiên của đồ án cơ sở. Trong tuần 9, nhóm quyết định hoàn thiện thêm "
             "hạng mục này trước khi bước vào tuần tổng kết cuối cùng, đồng thời cập nhật bản triển "
             "khai web hosting và chuẩn bị dữ liệu cho báo cáo tổng kết tuần 10.")
add_bullets(d9, [
    "Bổ sung chức năng \"Quên mật khẩu\" (đặt lại mật khẩu qua email) cho màn hình đăng nhập, dùng "
    "Firebase Authentication.",
    "Kiểm thử chức năng mới trên bản web đã triển khai và điều chỉnh giao diện theo góp ý thẩm mỹ.",
    "Cập nhật bản triển khai (Firebase Hosting) và đồng bộ mã nguồn lên kho Git từ xa.",
])

add_heading1(d9, "2. CÔNG VIỆC ĐÃ THỰC HIỆN")

add_heading2(d9, "2.1. Bổ sung chức năng Quên mật khẩu")
add_bullets(d9, [
    "Thêm phương thức AuthService.sendPasswordResetEmail() (lib/services/auth_service.dart) — gọi "
    "FirebaseAuth.sendPasswordResetEmail, tái sử dụng hàm ánh xạ lỗi tiếng Việt sẵn có "
    "(_mapErrorMessage) để báo lỗi rõ ràng (email không hợp lệ, không tìm thấy tài khoản...).",
    "Xây dựng màn hình mới ForgotPasswordScreen (lib/screens/forgot_password_screen.dart): nhập "
    "email, gửi liên kết đặt lại mật khẩu, chuyển sang trạng thái xác nhận \"đã gửi email\" kèm nút "
    "gửi lại — giữ đúng phong cách giao diện (màu #3B5BDB, bo góc 12px) của các màn hình Đăng nhập/"
    "Đăng ký sẵn có.",
    "Thêm liên kết \"Quên mật khẩu?\" vào màn hình Đăng nhập (lib/screens/login_screen.dart), điều "
    "hướng sang màn hình mới.",
])

add_heading2(d9, "2.2. Điều chỉnh giao diện theo góp ý")
add_bullets(d9, [
    "Bố trí ban đầu: liên kết \"Quên mật khẩu?\" đặt riêng, căn phải, phía trên nút Đăng nhập.",
    "Sau khi xem trên bản triển khai thật, bố trí này bị đánh giá là dàn trải, mất cân đối — nhóm "
    "chỉnh lại: gộp \"Quên mật khẩu?\" vào chung một hàng, căn giữa, cùng với liên kết \"Chưa có tài "
    "khoản? Đăng ký ngay\" đã có sẵn, phân cách bằng dấu \"•\" — vừa gọn vừa nhất quán với phong cách "
    "còn lại của màn hình.",
])
add_real_image_wide(d9, os.path.join(ASSETS9, "01_login_forgot_link.png"),
                     "Hình 1. Màn hình đăng nhập trên bản web đã triển khai — liên kết \"Quên mật "
                     "khẩu?\" đặt cùng hàng, căn giữa, cạnh liên kết đăng ký")
add_real_image_wide(d9, os.path.join(ASSETS9, "02_forgot_password_screen.png"),
                     "Hình 2. Màn hình Quên mật khẩu — nhập email để nhận liên kết đặt lại mật khẩu "
                     "qua Firebase Authentication")

add_heading2(d9, "2.3. Triển khai và kiểm thử trên bản web thật")
add_bullets(d9, [
    "Build lại bản web (flutter build web) và triển khai lên Firebase Hosting "
    "(https://english-ai-app-9b33e.web.app) bằng firebase deploy --only hosting; xác nhận lại tệp "
    "cấu hình .env vẫn được đóng gói đúng vào build/web/assets/.env (theo bản sửa lỗi đã ghi nhận ở "
    "tuần trước).",
    "Kiểm thử độc lập bản đã triển khai bằng trình duyệt Chromium headless (Playwright) chạy trên một "
    "phiên hoàn toàn mới (không cache, không service worker cũ) — xác nhận liên kết \"Quên mật khẩu?\" "
    "hiển thị đúng và điều hướng đúng sang màn hình đặt lại mật khẩu, độc lập với việc trình duyệt "
    "phía người dùng có thể đang giữ bản cache cũ.",
])

add_heading2(d9, "2.4. Đồng bộ mã nguồn")
add_bullets(d9, [
    "Commit các thay đổi liên quan đến chức năng Quên mật khẩu (auth_service.dart, login_screen.dart, "
    "forgot_password_screen.dart mới) thành một commit riêng và đẩy (push) lên nhánh main trên kho Git "
    "từ xa.",
])

add_heading1(d9, "3. KẾT QUẢ ĐẠT ĐƯỢC")
add_table(d9,
    ["Hạng mục", "Trạng thái", "Ghi chú"],
    [
        ["Chức năng Quên mật khẩu", "Hoàn thiện, đã kiểm thử", "Firebase Auth sendPasswordResetEmail"],
        ["Giao diện liên kết Quên mật khẩu trên màn hình đăng nhập", "Hoàn thiện, đã điều chỉnh theo góp ý thẩm mỹ", "Căn giữa, cùng hàng với liên kết đăng ký"],
        ["Triển khai web hosting", "Cập nhật thành công", "https://english-ai-app-9b33e.web.app"],
        ["Đồng bộ mã nguồn (Git)", "Hoàn thiện", "Nhánh main, đã push lên kho từ xa"],
    ])

add_heading1(d9, "4. KHÓ KHĂN & HƯỚNG GIẢI QUYẾT")
add_table(d9,
    ["Khó khăn", "Hướng giải quyết"],
    [
        ["Bố trí ban đầu của liên kết \"Quên mật khẩu?\" (căn phải, tách riêng phía trên nút Đăng "
         "nhập) không đẹp mắt khi xem trên bản triển khai thật",
         "Gộp chung một hàng, căn giữa với liên kết đăng ký sẵn có, dùng dấu phân cách \"•\" — vừa "
         "tiết kiệm không gian vừa đồng nhất bố cục với phần còn lại của màn hình."],
        ["Người dùng không thấy thay đổi ngay sau khi triển khai do trình duyệt/Service Worker của "
         "Flutter Web giữ bản cache cũ",
         "Hướng dẫn người dùng hard refresh (Ctrl+Shift+R) hoặc mở cửa sổ ẩn danh; đồng thời tự kiểm "
         "chứng độc lập bằng trình duyệt headless với phiên hoàn toàn mới để xác nhận bản deploy đúng, "
         "không phụ thuộc cache phía người dùng."],
    ])

add_heading1(d9, "5. KẾ HOẠCH TUẦN TIẾP THEO")
add_para(d9, "Tuần 10 là tuần cuối cùng của đồ án. Nhóm sẽ:")
add_bullets(d9, [
    "Tổng kết toàn diện toàn bộ đồ án sau 10 tuần thực hiện: rà soát và trình bày đầy đủ tất cả các "
    "chức năng đã xây dựng (từ các chức năng học viên cơ bản đến 5 tính năng AI nâng cao và toàn bộ "
    "khối quản trị), kèm minh hoạ hình ảnh thực tế.",
    "Cập nhật báo cáo tổng hợp đồ án và bộ slide thuyết trình bảo vệ với nội dung mới nhất (chức năng "
    "Quên mật khẩu).",
    "Rà soát lần cuối toàn bộ tài liệu, chuẩn bị sẵn sàng cho việc nộp và bảo vệ đồ án.",
])

add_heading1(d9, "6. KẾT LUẬN TUẦN 9")
add_para(d9, "Trong tuần 9, nhóm đã hoàn thiện chức năng Quên mật khẩu — hạng mục từng được ghi nhận "
             "là nằm ngoài phạm vi ưu tiên ở tuần 8 — bao gồm cả việc điều chỉnh giao diện theo góp ý "
             "thực tế, triển khai lên bản web và kiểm thử độc lập bằng trình duyệt headless. Toàn bộ "
             "thay đổi đã được đồng bộ lên kho mã nguồn. Đây là hạng mục kỹ thuật cuối cùng trước khi "
             "nhóm bước vào tuần 10 — tuần tổng kết toàn bộ đồ án.")

d9.save(r"C:\Users\Nam\Desktop\DACS\Bao_cao_tuan_9.docx")
print("Saved Bao_cao_tuan_9.docx")

# ---------------------------------------------------------------------------
# TUẦN 10 (BÁO CÁO CUỐI CÙNG — TỔNG KẾT TOÀN BỘ ĐỒ ÁN)
# ---------------------------------------------------------------------------
d10 = new_doc()
add_cover(d10, 10)

add_heading1(d10, "1. MỤC TIÊU TUẦN 10")
add_para(d10, "Tuần 10 là tuần cuối cùng của đồ án cơ sở ngành Công nghệ thông tin. Sau 9 tuần xây "
              "dựng, hoàn thiện và kiểm thử từng phần, mục tiêu của tuần này là tổng kết toàn diện "
              "toàn bộ hệ thống: trình bày đầy đủ tất cả các chức năng đã xây dựng kèm minh hoạ hình "
              "ảnh thực tế, xác nhận trạng thái hoàn thiện cuối cùng, và chuẩn bị sẵn sàng tài liệu "
              "cho việc nộp và bảo vệ đồ án.")
add_bullets(d10, [
    "Tổng hợp và trình bày đầy đủ toàn bộ chức năng của ứng dụng theo 4 nhóm: chức năng học viên cơ "
    "bản, 5 tính năng AI nâng cao, khối quản trị, và hạ tầng/triển khai.",
    "Chốt danh sách công nghệ, kiến trúc và quy mô mã nguồn cuối cùng của đồ án.",
    "Hoàn thiện báo cáo tổng hợp, bộ slide thuyết trình và video demo cho buổi bảo vệ.",
])

add_heading1(d10, "2. CÔNG VIỆC ĐÃ THỰC HIỆN")
add_para(d10, "Đây là phần tổng kết lại toàn bộ các chức năng đã xây dựng trong suốt quá trình thực "
              "hiện đồ án (tuần 1 đến tuần 9), trình bày theo 4 nhóm chức năng chính của hệ thống.")

add_heading2(d10, "2.1. Chức năng học viên cơ bản")
add_bullets(d10, [
    "Đăng ký, đăng nhập bằng email/mật khẩu (Firebase Authentication) và đặt lại mật khẩu qua email "
    "khi quên mật khẩu (hoàn thiện ở tuần 9).",
    "Xem danh sách bài học tiếng Anh, học từ vựng theo từng bài.",
    "Làm bài kiểm tra trắc nghiệm (quiz), tính điểm và lưu kết quả (QuizResultService).",
    "Theo dõi tiến độ học tập cá nhân theo thời gian (LearningProgressService).",
    "Nhận gợi ý học tập cá nhân hoá từ AI dựa trên kết quả làm bài (AiRecommendationService, tích hợp "
    "Gemini AI).",
])
add_real_image_wide(d10, os.path.join(ASSETS9, "01_login_forgot_link.png"),
                     "Hình 1. Màn hình đăng nhập (bản web) — có liên kết \"Quên mật khẩu?\"")
add_real_image_wide(d10, os.path.join(ASSETS9, "02_forgot_password_screen.png"),
                     "Hình 2. Màn hình đặt lại mật khẩu qua email")
add_real_image(d10, os.path.join(ASSETS7, "02_quiz.png"),
                "Hình 3. Làm bài kiểm tra trắc nghiệm trên thiết bị Android thật", height_cm=9)
add_real_image(d10, os.path.join(ASSETS7, "03_ai_recommendation.png"),
                "Hình 4. Gợi ý học tập cá nhân hoá từ AI", height_cm=9)
add_real_image(d10, os.path.join(ASSETS7, "05_progress.png"),
                "Hình 5. Theo dõi tiến độ học tập", height_cm=9)

add_heading2(d10, "2.2. Năm tính năng AI nâng cao")
add_bullets(d10, [
    "AI Camera từ vựng: chụp/chọn ảnh vật thể thực tế, Gemini Vision (gemini-3.5-flash-lite) nhận "
    "diện và trả về từ vựng tiếng Anh liên quan kèm nghĩa, phiên âm, lưu vào sổ từ vựng cá nhân.",
    "Từ điển AI: tra cứu nghĩa/phiên âm/từ loại/câu ví dụ cho từ hoặc cụm từ tiếng Anh bất kỳ (tái sử "
    "dụng chung service với AI Camera ở chế độ nhập văn bản).",
    "Luyện Nghe & Nói AI: xem video tiếng Anh nhúng (youtube_player_iframe), điền từ nghe được, ghi âm "
    "phát âm lại và nhận phân tích phát âm/nội dung từ AI (AiSpeakingFeedbackService).",
    "Luyện Shadowing: nghe câu mẫu qua Text-to-Speech, ghi âm nhại lại, được chấm điểm từng từ và nhận "
    "xét chi tiết từ \"AI Pronunciation Coach\" (ShadowingAiFeedbackService).",
    "Luyện Viết AI: nộp bài luận IELTS Task 1/Task 2, được AI chấm điểm theo 4 tiêu chí (Lexical "
    "Resource, Task Achievement, Grammatical Range, Coherence & Cohesion) với band điểm 0.0–9.0 "
    "(WritingAiGradingService).",
])
add_real_image(d10, os.path.join(ASSETS7, "06_shadowing.png"),
                "Hình 6. Luyện Shadowing — nghe câu mẫu và ghi âm nhại lại", height_cm=9)
add_real_image(d10, os.path.join(ASSETS7, "07_ai_dictionary.png"),
                "Hình 7. Kết quả tra cứu Từ điển AI", height_cm=9)
add_real_image(d10, os.path.join(ASSETS7, "08_listening_practice.png"),
                "Hình 8. Luyện Nghe & Nói AI — video nhúng và ô điền từ nghe được", height_cm=9)
add_real_image(d10, os.path.join(ASSETS7, "09_writing_result.png"),
                "Hình 9. Kết quả chấm điểm Luyện Viết AI theo 4 tiêu chí IELTS", height_cm=9)
add_para(d10, "Riêng tính năng AI Camera từ vựng sử dụng bộ chọn ảnh gốc của hệ điều hành (image "
              "picker), nằm ngoài khả năng chụp ảnh tự động của bộ kiểm thử integration_test — chức "
              "năng này được minh hoạ qua video demo đính kèm thay vì ảnh chụp màn hình tĩnh.")

add_heading2(d10, "2.3. Khối quản trị (Admin)")
add_bullets(d10, [
    "9 màn hình quản trị dạng CRUD: quản lý bài học, câu hỏi trắc nghiệm, từ vựng, bài nghe, bài "
    "Shadowing, bài viết, tài khoản người dùng, feature flags (bật/tắt tính năng), và trang thống kê "
    "tổng quan (dashboard).",
    "Phân quyền theo vai trò (student/admin) bằng Cloud Firestore Security Rules, áp dụng cho toàn bộ "
    "19 collection dữ liệu.",
])
add_real_image(d10, os.path.join(ASSETS7, "04_admin_dashboard.png"),
                "Hình 10. Trang quản trị hệ thống — thống kê tổng quan", height_cm=9)

add_heading2(d10, "2.4. Hạ tầng, công nghệ và triển khai")
add_bullets(d10, [
    "Nền tảng: Flutter (Dart, SDK ^3.12.2), chạy đa nền tảng Android và Web từ cùng một mã nguồn.",
    "Backend: Firebase (Authentication, Cloud Firestore, Security Rules) và Firebase Hosting cho bản "
    "web.",
    "AI: Google Gemini AI (nhận diện hình ảnh, phân tích văn bản, chấm điểm phát âm và bài viết).",
    "Các thư viện hỗ trợ chính: flutter_dotenv (biến môi trường), image_picker (chọn ảnh), "
    "youtube_player_iframe (video nhúng), record + audioplayers (ghi âm/phát âm thanh), flutter_tts "
    "(chuyển văn bản thành giọng nói).",
    "Giao diện dark theme áp dụng đồng bộ trên toàn bộ ứng dụng, đã kiểm thử trên thiết bị Android "
    "thật.",
    "Quy mô mã nguồn cuối cùng: 18 model, 19 service, 31 màn hình (22 màn hình học viên + 9 màn hình "
    "quản trị), tương ứng 19 collection trên Cloud Firestore.",
    "Bản web đã triển khai công khai tại https://english-ai-app-9b33e.web.app (Firebase Hosting).",
])

add_heading2(d10, "2.5. Tài liệu và chuẩn bị bảo vệ đồ án")
add_bullets(d10, [
    "Hoàn thành đầy đủ 10 báo cáo tuần theo đúng tiến độ, ghi nhận trung thực quá trình xây dựng, "
    "kiểm thử và các điều chỉnh trong suốt đồ án.",
    "Báo cáo tổng hợp đồ án (thuyết minh) tổng hợp mục tiêu, phạm vi, công nghệ, kiến trúc và kết quả "
    "của toàn bộ quá trình thực hiện.",
    "Bộ slide thuyết trình bảo vệ đồ án và video demo minh hoạ trực tiếp các chức năng chính.",
])

add_heading1(d10, "3. KẾT QUẢ ĐẠT ĐƯỢC")
add_para(d10, "Trạng thái hoàn thiện cuối cùng của toàn bộ đồ án:")
add_table(d10,
    ["Nhóm chức năng", "Trạng thái", "Ghi chú"],
    [
        ["Chức năng học viên cơ bản (đăng ký/đăng nhập/quên mật khẩu, bài học, từ vựng, quiz, tiến độ, gợi ý AI)", "Hoàn thiện, đã kiểm thử", ""],
        ["5 tính năng AI nâng cao (Camera, Từ điển, Nghe & Nói, Shadowing, Viết)", "Hoàn thiện, đã kiểm thử trên Android & Web", ""],
        ["Khối quản trị (9 màn hình CRUD, phân quyền, feature flags, dashboard)", "Hoàn thiện, đã kiểm thử", ""],
        ["Giao diện dark theme", "Hoàn thiện, đồng bộ toàn ứng dụng", ""],
        ["Triển khai web hosting", "Hoàn thiện, đang hoạt động", "https://english-ai-app-9b33e.web.app"],
        ["Báo cáo tuần (10 tuần)", "Hoàn thiện", "File .docx"],
        ["Báo cáo tổng hợp đồ án", "Hoàn thiện", "File .docx"],
        ["Slide thuyết trình bảo vệ", "Hoàn thiện", "File .pptx"],
        ["Video demo minh hoạ chức năng", "Hoàn thiện", "Đính kèm cùng báo cáo/slide"],
    ])

add_heading1(d10, "4. KHÓ KHĂN & HƯỚNG GIẢI QUYẾT")
add_table(d10,
    ["Khó khăn", "Hướng giải quyết"],
    [
        ["Chưa có thiết bị/máy Mac để build và kiểm thử trên iOS trong suốt quá trình thực hiện",
         "Ưu tiên hoàn thiện đầy đủ trên Android và Web — hai nền tảng đã kiểm thử kỹ lưỡng và ổn "
         "định; ghi nhận rõ trong tài liệu để bổ sung iOS khi có điều kiện sau đồ án."],
        ["Các tính năng AI phụ thuộc dịch vụ Gemini bên thứ ba (giới hạn quota bản miễn phí, từng phải "
         "đổi model do bị deprecate giữa đồ án)",
         "Tách riêng lớp service gọi AI (mỗi tính năng một service độc lập) để việc đổi model chỉ ảnh "
         "hưởng cục bộ; theo dõi thông báo từ Google AI Studio để cập nhật kịp thời."],
        ["Hành vi ghi âm/truy cập microphone khác nhau giữa Android và trình duyệt web",
         "Tách riêng luồng xử lý theo nền tảng (kIsWeb) cho các tính năng liên quan đến ghi âm (Nghe & "
         "Nói AI, Shadowing); kiểm thử lại trên cả 2 nền tảng mỗi khi nâng cấp package ghi âm."],
        ["AI Camera từ vựng dùng bộ chọn ảnh gốc của hệ điều hành, không thể chụp ảnh tự động qua "
         "integration_test",
         "Minh hoạ chức năng này bằng video demo thực tế thay vì ảnh chụp màn hình tĩnh."],
    ])

add_heading1(d10, "5. KẾ HOẠCH CHUẨN BỊ BẢO VỆ ĐỒ ÁN")
add_bullets(d10, [
    "Rà soát lần cuối toàn bộ 10 báo cáo tuần, báo cáo tổng hợp và bộ slide trước khi nộp.",
    "Luyện tập trình bày nội dung đồ án và thao tác demo trực tiếp cho buổi bảo vệ.",
    "Chuẩn bị sẵn sàng trả lời các câu hỏi liên quan đến kiến trúc hệ thống, lựa chọn công nghệ và "
    "hướng phát triển tiếp theo của ứng dụng.",
])

add_heading1(d10, "6. KẾT LUẬN TUẦN 10 VÀ TỔNG KẾT ĐỒ ÁN")
add_para(d10, "Sau 10 tuần thực hiện, nhóm đã hoàn thiện ứng dụng di động học tiếng Anh dựa trên AI "
              "với đầy đủ các chức năng học viên cơ bản, 5 tính năng AI nâng cao (AI Camera từ vựng, "
              "Từ điển AI, Luyện Nghe & Nói AI, Luyện Shadowing, Luyện Viết AI), khối quản trị hoàn "
              "chỉnh và bản triển khai web đang hoạt động công khai. Toàn bộ hệ thống đã được kiểm "
              "thử kỹ lưỡng trên cả Android và trình duyệt web, cùng bộ tài liệu báo cáo, slide và "
              "video demo đầy đủ, sẵn sàng cho việc nộp và bảo vệ đồ án.")

d10.save(r"C:\Users\Nam\Desktop\DACS\Bao_cao_tuan_10.docx")
print("Saved Bao_cao_tuan_10.docx")
